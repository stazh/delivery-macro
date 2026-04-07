import config
import os
import re
from tkinter import filedialog, messagebox, Tk
from typing import List, Dict, Any
from app import root

# Versuche, win32com.client zu importieren
try:
    import win32com.client as win32
except ImportError as e:
    messagebox.showerror("Fehler", f"Fehler beim Importieren von win32com.client: {e}", parent=root)
    raise

# Versuche, openpyxl zu importieren
try:
    import openpyxl
except ImportError as e:
    messagebox.showerror("Fehler", f"Fehler beim Importieren von openpyxl: {e}", parent=root)
    raise

# Versuche, xlwings zu importieren
try:
    import xlwings as xw
except ImportError as e:
    messagebox.showerror("Fehler", f"Fehler beim Importieren von xlwings: {e}", parent=root)
    raise


def import_delivery_list_word() -> None:
    """Importiert eine Ablieferungsverzeichnis aus einer Word-Datei."""
    root = Tk()
    root.withdraw()

    # Öffnen des Datei-Dialogs zur Auswahl einer Word-Datei
    file = filedialog.askopenfilename(filetypes=[("Word-Dokument", "*.docx")], title="Bitte Word-Datei auswählen...")
    if not file:
        messagebox.showinfo("Abbruch", "Kein Dateipfad ausgewählt. Der Vorgang wird abgebrochen.", parent=root)
        return

    doc_props, table_data = read_docx_data(file)
    create_delivery_list(doc_props, table_data)


def import_delivery_list_excel() -> None:
    """Importiert eine Ablieferungsverzeichnis aus einer Excel-Datei."""
    root = Tk()
    root.withdraw()

    try:
        app = xw.apps.active
        app.screen_updating
        workbook = app.books[config.DELIVERY_MACRO_FILE_NAME]
    except Exception as e:
        messagebox.showerror("Fehler", "Excel‑Instanz konnte nicht gefunden werden oder Datei ist nicht offen.", parent=root)
        return

    try:
        app.screen_updating = True
        doc_props, table_data = read_excel_data(workbook)
        create_delivery_list(doc_props, table_data)
    except Exception as e:
        messagebox.showerror("Fehler", f"Fehler beim Verarbeiten der Abliefererdaten: {e}", parent=root)


def create_delivery_list(doc_props: Dict[str, Any], table_data: List[List[str]]) -> None:
    """Erstellt die Ablieferungsverzeichnis in einer Excel-Datei."""
    try:
        app = xw.apps.active
        app.screen_updating = False
        stammdaten_workbook = app.books.open(config.DATA_FILE_PATH)
        dateipfade_sheet = stammdaten_workbook.sheets[config.DATA_PATH_SHEET_NAME]
        template_rel = dateipfade_sheet.range(config.DATA_PATH_FIELD_DELIVERY_LIST).value
        stammdaten_workbook.close()
        exc_file = os.path.abspath(template_rel)

        if not os.path.exists(exc_file):
            messagebox.showerror("Fehler", f"Die Vorlagen‑Datei wurde nicht gefunden:\n{exc_file}", parent=root)
            return

        # Zielordner für das Ablieferungsverzeichnis aus config lesen und sicherstellen, dass der Ordner existiert
        output_folder = config.CREATED_DOCUMENT_FOLDER_PATH
        if not os.path.exists(output_folder):
            os.makedirs(output_folder)

        # Der Name der Ablieferungsverzeichnis-Datei aus config und sicherstellen, dass die Datei nicht bereits existiert
        document_number = 1
        exc_file_path = os.path.join(output_folder, f"{document_number:02d}_{config.DELIVERY_LIST_FILE_NAME}")
        while os.path.isfile(exc_file_path):
            document_number += 1
            exc_file_path = os.path.join(output_folder, f"{document_number:02d}_{config.DELIVERY_LIST_FILE_NAME}")
        
        # Umwandlung des Zielpfades in einen absoluten Pfad
        exc_file_path = os.path.abspath(exc_file_path)

        # Lade die Vorlage und speichere sie im Zielordner
        wb = openpyxl.load_workbook(exc_file)
        wb.save(exc_file_path)
        
        # Öffne die neu gespeicherte Excel-Datei
        app.screen_updating = True
        workbook = app.books.open(exc_file_path)
        ws = workbook.sheets[config.DELIVERY_LIST_SHEET_NAME]
        
        # Ersetze Platzhalter in der Excel-Datei
        replace_placeholders_in_excel(ws, doc_props)
        insert_table_data(ws, table_data)
        
        # Speichern der Excel-Datei
        workbook.save()

        messagebox.showinfo("Fertig", f"Das Ablieferungsverzeichnis wurde erfolgreich erstellt", parent=root)
    except Exception as e:
        messagebox.showerror("Fehler", f"Fehler beim Erstellen der Excel‑Datei: {e}", parent=root)


def read_docx_data(file_path: str) -> tuple[Dict[str, Any], List[List[str]]]:
    """Liest die Daten aus einer Word-Datei."""
    word = win32.DispatchEx("Word.Application")
    word.visible = False

    try:
        doc = word.Documents.Open(os.path.abspath(file_path))

        table_data = []
        raw_data = []
        doc = word.ActiveDocument
        table = doc.Tables[2]
        declinedTable = doc.Tables[3]
        ang_lfm = 0
        ang_gb = 0

        for i in range(2, table.Rows.Count + 1):
            row = table.Rows(i)
            first_column_data = row.Cells(1).Range.Text.strip()
            
            # Prüfen, ob die Zeile nur eine Zelle hat
            if row.Cells.Count == 1:
                third_column_data = ""
            else:
                third_column_data = row.Cells(3).Range.Text.strip()
            
            raw_data.append([first_column_data, third_column_data])

            # Lfm und GB aus der ersten Spalte extrahieren
            lfm_matches = list(re.finditer(config.REGEX_WORD_DELIVERY_AGREEMENT_SEARCH_LFM, first_column_data))
            gb_matches = list(re.finditer(config.REGEX_WORD_DELIVERY_AGREEMENT_SEARCH_GB, first_column_data))

            if lfm_matches:
                lfm_value = lfm_matches[-1].group(0).replace(" Lfm", "")
                ang_lfm += round(float(lfm_value), 2)

            if gb_matches:
                gb_value = gb_matches[-1].group(0).replace(" GB", "")
                ang_gb += round(float(gb_value), 2)

        for i in range(2, declinedTable.Rows.Count + 1):
            row = declinedTable.Rows(i)
            second_column_data = row.Cells(2).Range.Text.strip()

            # Lfm und GB aus der zweiten Spalte extrahieren
            lfm_match = re.search(rf'{config.REGEX_WORD_DELIVERY_AGREEMENT_SEARCH_LFM}', second_column_data)
            gb_match = re.search(rf'{config.REGEX_WORD_DELIVERY_AGREEMENT_SEARCH_GB}', second_column_data)

            if lfm_match:
                lfm_value = lfm_match.group(0).replace(" Lfm", "")
                ang_lfm += round(float(lfm_value), 2)

            if gb_match:
                gb_value = gb_match.group(0).replace(" GB", "")
                ang_gb += round(float(gb_value), 2)

        # Dokument-Eigenschaften zusammenstellen
        doc_props = {
            "DirName": config.DIR_NAME_PLACEHOLDER,
            "AmtZeichen": config.AMT_ZEICHEN_PLACEHOLDER,
            "AmtPName": config.AMT_P_NAME_PLACEHOLDER,
            "ÜbernDatum": config.UEBERN_DATUM_PLACEHOLDER,
            "AngLfm": ang_lfm,
            "AngGB": ang_gb,
            "ÜbernLfm": "",
            "ÜbernGB": "",
            "AblNummer": config.ABL_NUMMER_PLACEHOLDER
        }

        table_data = convert_word_table_data(raw_data)
        word.Quit()
    except Exception as e:
        word.Quit()
        messagebox.showerror("Fehler", f"Fehler beim Lesen der Word-Datei: {e}", parent=root)

    return doc_props, table_data


def convert_word_table_data(raw_data: List[List[str]]) -> List[List[str]]:
    """Konvertiert die rohen Daten der Word-Tabelle in das benötigte Format für Excel."""
    table_data = []
    for row in raw_data:
        inhalt = row[0]
        zeitraum = ""
        lfm = "0"
        gb = "0"
        medium = ""

        # Jahr und Medium extrahieren
        year_matches = list(re.finditer(config.REGEX_WORD_DELIVERY_AGREEMENT_SEARCH_YEAR, row[0]))

        if year_matches:
            last_year_match = year_matches[-1]

            year_index = last_year_match.start()
            parts = row[0][:year_index].strip(), row[0][year_index:].strip()

            inhalt = parts[0].removesuffix(',')
            zeitraum = parts[1].split(',')[1].strip() if ',' in parts[1] else parts[1].strip()

            medium_match = re.search(config.REGEX_WORD_DELIVERY_AGREEMENT_SEARCH_MEDIUM, row[0])
            if medium_match:
                medium = medium_match.group(1).strip()

        if zeitraum == '':
            table_data.append([inhalt.strip().replace('\n', '').replace('\r', '').replace('\x07', ''), '', '', '', ''])
            continue

        # Bereinige den Text in row[1]
        row[1] = row[1].strip().replace('\n', '').replace('\r', '').replace('\x07', '')

        if "(" in row[1] and ")" in row[1]:
            parts = row[1].split('(')
            number_part = parts[0].strip()
            medium_part = f"({parts[1].split(')')[0]})".strip()
        else:
            number_part = row[1]
            medium_part = ""

        lfm = config.DEFAULT_VALUE_LFM_GB
        gb = config.DEFAULT_VALUE_LFM_GB

        if '/' in number_part:
            parts = [p.strip() for p in number_part.split('/')]
            lfm = parts[0].replace(" Lfm", "").strip()
            gb = parts[1].replace(" GB", "").strip()
        else:
            if "Lfm" in number_part:
                lfm = number_part.replace(" Lfm", "").strip() if number_part.replace(" Lfm", "").strip() != '0.0' else config.DEFAULT_VALUE_LFM_GB
                gb = "0"
            elif "GB" in number_part:
                gb = number_part.replace(" GB", "").strip() if number_part.replace(" GB", "").strip() != '0.0' else config.DEFAULT_VALUE_LFM_GB
                lfm = "0"
            else:
                lfm = "0"
                gb = "0"

        if lfm == config.DEFAULT_VALUE_LFM_GB and gb == config.DEFAULT_VALUE_LFM_GB:
            continue

        table_data.append([inhalt, zeitraum, lfm, gb, medium_part.replace('(', '').replace(')', '')])
    return table_data


def read_excel_data(workbook) -> tuple[Dict[str, Any], List[List[str]]]:
    """Liest die Ablieferungsdaten aus einer Excel-Datei."""
    ws = workbook.sheets[config.CURRENT_OFFER_SHEET_NAME]

    doc_props = {
        "DirName": ws.range(config.CURRENT_OFFER_FIELD_DIR_NAME).value or "",
        "AmtZeichen": ws.range(config.CURRENT_OFFER_FIELD_AMT_ZEICHEN).value or "",
        "AmtPName": ws.range(config.CURRENT_OFFER_FIELD_AMT_P_NAME).value or "",
        "ÜbernDatum": ws.range(config.CURRENT_OFFER_FIELD_UEBERN_DATUM).value or config.UEBERN_DATUM_PLACEHOLDER,
        "AngLfm": "0" if ws.range(config.CURRENT_OFFER_FIELD_ANG_LFM).value == 0.0 else str(ws.range(config.CURRENT_OFFER_FIELD_ANG_LFM).value),
        "ÜbernLfm": "0" if ws.range(config.CURRENT_OFFER_FIELD_UEBERN_LFM).value == 0.0 else str(ws.range(config.CURRENT_OFFER_FIELD_UEBERN_LFM).value),
        "AngGB": "0" if ws.range(config.CURRENT_OFFER_FIELD_ANG_GB).value == 0.0 else str(ws.range(config.CURRENT_OFFER_FIELD_ANG_GB).value),
        "ÜbernGB": "0" if ws.range(config.CURRENT_OFFER_FIELD_UEBERN_GB).value == 0.0 else str(ws.range(config.CURRENT_OFFER_FIELD_UEBERN_GB).value),
        "AblNummer": ws.range(config.CURRENT_OFFER_FIELD_ABL_NUMMER).value or ""
    }

    table_data = []
    rows = ws.range(f"A{2}:E{ws.cells.last_cell.row}").value
    reason_row = ws.range(f"G{2}:G{ws.cells.last_cell.row}").value
    accepted_data_row = ws.range(f"J{2}:L{ws.cells.last_cell.row}").value
    if rows:
        for row, accepted_row, reason in zip(rows, accepted_data_row, reason_row):
            if all(cell is None or str(cell).strip() == "" for cell in row) or (reason == config.SELECTION_REASON_DECLINED_TEXT) or (reason == config.SELECTION_REASON_NO_TAKEOVER_TEXT):
                continue

            lfm_accepted = accepted_row[0] if accepted_row[0] is not None else "0"
            gb_accepted = accepted_row[1] if accepted_row[1] is not None else "0"
            medium_accepted = accepted_row[2] if accepted_row[2] is not None else ""

            if lfm_accepted == "0" and gb_accepted == "0":
                lfm_accepted = ''
                gb_accepted = ''

            row[2] = lfm_accepted
            row[3] = gb_accepted
            row[4] = medium_accepted
            table_data.append(row)
    return doc_props, table_data


def replace_placeholders_in_excel(ws, doc_props: Dict[str, Any]) -> None:
    """Ersetzt Platzhalter in einer Excel-Tabelle durch tatsächliche Werte."""
    placeholders = [
        config.DIR_NAME_PLACEHOLDER, config.AMT_ZEICHEN_PLACEHOLDER, config.AMT_P_NAME_PLACEHOLDER, config.UEBERN_DATUM_PLACEHOLDER,
        config.ANG_LFM_PLACEHOLDER, config.ANG_GB_PLACEHOLDER, config.UEBERN_LFM_PLACEHOLDER, config.UEBERN_GB_PLACEHOLDER, config.ABL_NUMMER_PLACEHOLDER
    ]
    
    values = [
        str(doc_props["DirName"]),
        str(doc_props["AmtZeichen"]),
        str(doc_props["AmtPName"]),
        str(doc_props["ÜbernDatum"]),
        str(doc_props["AngLfm"]),
        str(doc_props["AngGB"]),
        str(doc_props["ÜbernLfm"]),
        str(doc_props["ÜbernGB"]),
        str(doc_props["AblNummer"])
    ]

    for placeholder, value in zip(placeholders, values):
        for row in ws.range(config.DELIVERY_LIST_PLACEHOLDER_RANGE).rows:
            for cell in row:
                if cell.value and isinstance(cell.value, str):
                    if placeholder in cell.value:
                        cell.value = cell.value.replace(placeholder, value)


def insert_table_data(ws, table_data: List[List[str]]) -> None:
    """Fügt die Tabellen-Daten in die Excel-Datei ein."""
    row_start = config.DELIVERY_LIST_START_DATA_ROW
    for row in table_data:
        cell_range = ws.range(f'B{row_start}:F{row_start}')
        cell_range.value = row
        row_start += 1
